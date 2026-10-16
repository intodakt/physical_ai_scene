"""S7 metrics and the experiment matrix (project guide Section 24).

    python -m src.evaluation.metrics                       # mock sequence, configs/s7_eval_queries.json
    python -m src.evaluation.metrics --frames <S1 frames> --queries <file> --out outputs/evaluation

Writes outputs/evaluation/experiment_matrix.xlsx (+ CSV copies and metrics.json):
  * pipeline latency per stage and FPS (S1 -> S6);
  * sim-to-real success: validated ActionGoals the virtual agent reached without entering a zone;
  * hallucination catch rate: wrong LLM answers that the S5 verifier overrode;
  * answer accuracy, unsafe-goal rejection and tracking ID switches.
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from src.integration.pipeline import STAGES

DEFAULT_QUERIES = Path("configs/s7_eval_queries.json")


def _label_of(target_id):
    return None if target_id is None else target_id.rsplit("_", 1)[0]


def evaluate_queries(pipe, queries: list[dict]) -> list[dict]:
    rows = []
    for q in queries:
        r = pipe.ask(q["query"])
        rr, goal, sim = r["reasoning_result"], r["action_goal"], r["simulation"]
        got = _label_of(rr["target_id"])
        claim, expect = q.get("llm_claim"), q.get("expect_target")
        claim_wrong = claim is not None and claim != expect
        rows.append({
            "query": q["query"], "expected": expect, "answer_id": rr["target_id"],
            "verified": rr["verification_status"], "correct": got == expect,
            "llm_claim": claim, "llm_claim_wrong": claim_wrong,
            "hallucination_caught": claim_wrong and got != claim,
            "validated": goal["validated"], "expect_validated": q.get("expect_validated"),
            "goal_reason": goal["reason"], "sim_success": sim["success"],
            "sim_final_error_m": sim["final_error_m"], "s5_ms": r["timings_ms"]["s5_reasoning"],
            "s6_ms": r["timings_ms"]["s6_action"], "final_answer": rr["final_answer"],
        })
    return rows


def id_switches(frames: list[dict]) -> int:
    """Objects that received more than one persistent ID (same label) over the sequence."""
    ids = {}
    for f in frames:
        for pid in f["ids"].values():
            ids.setdefault(_label_of(pid), set()).add(pid)
    return sum(len(v) - 1 for v in ids.values())


def _rate(num, den):
    return round(num / den, 3) if den else None


def summarize(summary: dict, rows: list[dict]) -> dict:
    validated = [r for r in rows if r["validated"]]
    wrong_claims = [r for r in rows if r["llm_claim_wrong"]]
    unsafe = [r for r in rows if r["expect_validated"] is False and r["expected"] is not None]
    return {
        "frames": summary["n_frames"],
        "s3_mode": summary["s3_mode"],
        "s5_parser": summary["s5_parser"],
        "mean_latency_ms_S1_to_S6": summary["mean_latency_ms"],
        "fps": summary["fps"],
        **{f"latency_ms_{s}": v for s, v in summary["stage_latency_ms"].items()},
        "answer_accuracy": _rate(sum(r["correct"] for r in rows), len(rows)),
        "hallucination_catch_rate": _rate(sum(r["hallucination_caught"] for r in wrong_claims), len(wrong_claims)),
        "wrong_llm_claims": len(wrong_claims),
        "sim_to_real_success": _rate(sum(r["sim_success"] for r in validated), len(validated)),
        "goals_validated": len(validated),
        "unsafe_goals_rejected": _rate(sum(not r["validated"] for r in unsafe), len(unsafe)),
        "goal_decisions_correct": _rate(sum(r["validated"] == r["expect_validated"] for r in rows), len(rows)),
        "tracking_id_switches": id_switches(summary["frames"]),
        "contract_errors": len(summary["contract_errors"]),
    }


def write_matrix(out_dir, metrics: dict, rows: list[dict], frames: list[dict]) -> Path:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    frame_rows = [{"frame": f["frame"], "objects": f["n_objects"], "rejected": f["n_rejected"],
                   **{f"{k}_ms": v for k, v in f["timings_ms"].items()},
                   "persistent_ids": " ".join(sorted(set(f["ids"].values())))} for f in frames]
    sheets = {"summary": [{"metric": k, "value": v} for k, v in metrics.items()],
              "queries": rows, "latency_per_frame": frame_rows}
    for name, data in sheets.items():
        with open(out / f"{name}.csv", "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(data[0]))
            w.writeheader()
            w.writerows(data)
    (out / "metrics.json").write_text(json.dumps({"summary": metrics, "queries": rows}, indent=2), encoding="utf-8")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return out / "summary.csv"
    wb = Workbook()
    wb.remove(wb.active)
    for name, data in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(list(data[0]))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in data:
            ws.append([v if isinstance(v, (int, float, str, bool)) or v is None else str(v) for v in row.values()])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) for c in col)))
        ws.freeze_panes = "A2"
    path = out / "experiment_matrix.xlsx"
    wb.save(path)
    return path


def main():
    from main_pipeline import run

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--frames", default=None)
    ap.add_argument("--bag", default=None)
    ap.add_argument("--queries", default=str(DEFAULT_QUERIES))
    ap.add_argument("--parser", choices=["auto", "ollama", "rules"], default=None)
    ap.add_argument("--out", default="outputs/evaluation")
    a = ap.parse_args()
    pipe, summary = run(a.frames, a.bag, out=str(Path(a.out) / "run"), parser=a.parser, quiet=True)
    rows = evaluate_queries(pipe, json.loads(Path(a.queries).read_text(encoding="utf-8"))["queries"])
    metrics = summarize(summary, rows)
    path = write_matrix(a.out, metrics, rows, summary["frames"])
    for k, v in metrics.items():
        print(f"  {k:<32} {v}")
    for r in rows:
        mark = "ok " if r["correct"] and r["validated"] == r["expect_validated"] else "BAD"
        print(f"  [{mark}] {r['query']:<42} -> {r['answer_id']}  validated={r['validated']}")
    print("->", path)


if __name__ == "__main__":
    main()
